import io
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

MIN_FONT_PT = 10


@dataclass
class Issue:
    severity: str  # "high" or "medium"
    message: str


def check(data: bytes, filename: str) -> list[Issue]:
    """Structural ATS risks. Only DOCX exposes the layout needed for this."""
    if Path(filename).suffix.lower() != ".docx":
        return []

    try:
        doc = Document(io.BytesIO(data))
    except Exception:
        return []

    issues: list[Issue] = []

    if doc.tables:
        issues.append(
            Issue("high", f"Contains {len(doc.tables)} table(s). Many parsers read table cells out of order.")
        )

    if doc.inline_shapes:
        issues.append(Issue("high", "Contains images. Text inside an image is invisible to parsers."))

    if "txbxContent" in doc.element.xml:
        issues.append(Issue("high", "Contains text boxes. Their contents are often skipped entirely."))

    if any(_column_count(section) > 1 for section in doc.sections):
        issues.append(Issue("high", "Uses multiple columns. Parsers commonly interleave the text."))

    if _has_header_or_footer_text(doc):
        issues.append(Issue("medium", "Text in headers or footers is frequently dropped. Keep contact details in the body."))

    smallest = _smallest_font(doc)
    if smallest is not None and smallest < MIN_FONT_PT:
        issues.append(
            Issue("medium", f"Font as small as {smallest:g}pt. Keep body text at {MIN_FONT_PT}pt or above.")
        )

    return issues


def _column_count(section) -> int:
    cols = section._sectPr.find(qn("w:cols"))
    if cols is None:
        return 1
    return int(cols.get(qn("w:num"), "1"))


def _has_header_or_footer_text(doc: Document) -> bool:
    for section in doc.sections:
        for part in (section.header, section.footer):
            if any(p.text.strip() for p in part.paragraphs):
                return True
    return False


def _smallest_font(doc: Document) -> float | None:
    sizes = [
        run.font.size.pt
        for paragraph in doc.paragraphs
        for run in paragraph.runs
        if run.font.size is not None
    ]
    return min(sizes) if sizes else None



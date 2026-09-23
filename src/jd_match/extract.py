import io
from pathlib import Path

from docx import Document
from pypdf import PdfReader

MAX_BYTES = 5 * 1024 * 1024
SUPPORTED = {".txt", ".md", ".docx", ".pdf"}


class ExtractionError(Exception):
    """Raised when a file cannot be read as text."""


def from_path(path: Path) -> str:
    if not path.exists():
        raise ExtractionError(f"File not found: {path}")
    return from_bytes(path.read_bytes(), path.name)


def from_bytes(data: bytes, filename: str) -> str:
    suffix = Path(filename).suffix.lower()

    if suffix not in SUPPORTED:
        raise ExtractionError(f"Unsupported file type: {suffix or filename}")
    if len(data) > MAX_BYTES:
        raise ExtractionError("File is larger than the 5 MB limit.")

    if suffix in {".txt", ".md"}:
        return data.decode("utf-8", errors="replace")
    if suffix == ".docx":
        return _from_docx(data)
    return _from_pdf(data)


def _from_docx(data: bytes) -> str:
    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError("Could not read the DOCX file.") from exc

    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)

    return "\n".join(parts)


def _from_pdf(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as exc:
        raise ExtractionError("Could not read the PDF file.") from exc


